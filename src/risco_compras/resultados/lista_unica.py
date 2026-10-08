"""Lista única de casos: junta os alertas das regras, gera a planilha de análise e registra os desfechos."""

from __future__ import annotations

import argparse
import math
import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd

from ..config import RAIZ, carregar_parametros

SITUACOES = ["Pendente", "Em análise", "Procedente", "Improcedente", "Fora de escopo"]
MOTIVOS = ["Erro de cadastro ou de dado", "Situação regular ou prevista na política", "Vínculo não confirmado",
           "Valor irrelevante", "Já tratado em outro caso", "Outro"]
PROVIDENCIAS = ["Nenhuma", "Recomendação ao departamento", "Auditoria aberta", "Bloqueio do fornecedor",
                "Encaminhado ao jurídico", "Ressarcimento"]
COLUNAS_DESFECHO = ["situacao", "motivo", "providencia", "analista", "data_analise", "observacao"]
TITULOS_DESFECHO = {"situacao": "Situação", "motivo": "Motivo (se improcedente)",
                    "providencia": "Providência (se procedente)", "analista": "Analista",
                    "data_analise": "Data da análise", "observacao": "Observação"}

PASTA_PADRAO = RAIZ / "dados" / "resultados"
_ARQUIVO_ALERTAS = re.compile(r"^alertas_(?P<regra>R\w*?)_(?P<carimbo>\d{8}_\d{4,6})\.csv$")
_PAR_VINCULADO = re.compile(r"Vínculo com (\d{14})")
REGRA_DE_PARES = "R3"


def pastas(pasta: Path | None = None) -> dict[str, Path]:
    base = Path(pasta) if pasta else PASTA_PADRAO
    return {"base": base, "desfechos": base / "desfechos", "painel": base / "painel"}


def config_lote(parametros: dict | None = None) -> dict:
    conf = (parametros or carregar_parametros()).get("lista_unica", {})
    return {"nota_1": tuple(conf.get("regras_nota_1", [])), "sempre": tuple(conf.get("regras_sempre", [])),
            "maiores_valores": int(conf.get("maiores_valores", 10))}


def regras_ativas() -> dict[str, dict]:
    return {rid: conf for rid, conf in carregar_parametros().get("regras", {}).items() if conf.get("ativa", True)}


def ultimos_arquivos(pasta: Path, ids: list[str]) -> dict[str, Path]:
    """Arquivo mais recente de cada regra (pelo carimbo no nome)."""
    achados: dict[str, tuple[str, Path]] = {}
    for arquivo in Path(pasta).glob("alertas_*.csv"):
        m = _ARQUIVO_ALERTAS.match(arquivo.name)
        if not m or m["regra"] not in ids:
            continue
        if m["regra"] not in achados or m["carimbo"] > achados[m["regra"]][0]:
            achados[m["regra"]] = (m["carimbo"], arquivo)
    return {rid: caminho for rid, (_, caminho) in sorted(achados.items())}


def _so_digitos(valor) -> str:
    texto = "" if valor is None or (isinstance(valor, float) and math.isnan(valor)) else str(valor)
    return re.sub(r"\D", "", re.sub(r"\.0$", "", texto.strip()))


def carregar_alertas(arquivos: dict[str, Path], nomes_fornecedores: dict[str, str] | None = None) -> pd.DataFrame:
    tabelas = []
    for rid, arquivo in arquivos.items():
        df = pd.read_csv(arquivo, sep=";", dtype=str, encoding="utf-8").fillna("")
        df["id_regra"] = rid
        df["arquivo"] = arquivo.name
        tabelas.append(df)
    if not tabelas:
        return pd.DataFrame()
    alertas = pd.concat(tabelas, ignore_index=True)
    alertas["pontuacao"] = pd.to_numeric(alertas["pontuacao"], errors="coerce").fillna(0.0).clip(0, 1)
    alertas["valor_envolvido"] = pd.to_numeric(alertas["valor_envolvido"], errors="coerce").fillna(0.0)
    alertas["documento"] = alertas["cnpj_fornecedor"].map(_so_digitos).str.zfill(14)
    alertas["classe"] = alertas["evidencia"].str.extract(r"^\[([^\]]+)\]", expand=False).fillna("")
    alertas["espelho"] = False
    alertas = espelhar_pares(alertas)
    alertas["fornecedor"] = alertas["documento"].map(nomes_fornecedores or {}).fillna("")
    alertas["chave_caso"] = alertas["documento"] + "|" + alertas["departamento"]
    return alertas


def espelhar_pares(alertas: pd.DataFrame) -> pd.DataFrame:
    """Duplica o alerta da regra de pares para a segunda empresa do par (espelho=True)."""
    pares = alertas[(alertas["id_regra"] == REGRA_DE_PARES) & ~alertas["espelho"]]
    linhas = []
    for _, alerta in pares.iterrows():
        m = _PAR_VINCULADO.search(alerta["evidencia"])
        if not m or m.group(1) == alerta["documento"]:
            continue
        outra = alerta.copy()
        outra["documento"] = outra["cnpj_fornecedor"] = m.group(1)
        outra["evidencia"] = (alerta["evidencia"].replace(m.group(0), f"Vínculo com {alerta['documento']}", 1)
                              + " (espelho: a regra registrou o par na outra empresa)")
        outra["espelho"] = True
        linhas.append(outra)
    if not linhas:
        return alertas
    return pd.concat([alertas, pd.DataFrame(linhas)], ignore_index=True)


def nota_combinada(notas_por_regra: list[float]) -> float:
    """1 − ∏(1 − nota): cada regra é um indício independente; mais regras, nota maior (máx. 1)."""
    produto = 1.0
    for nota in notas_por_regra:
        produto *= 1.0 - max(0.0, min(1.0, float(nota)))
    return round(1.0 - produto, 3)


def ler_desfechos(pasta_desfechos: Path) -> pd.DataFrame:
    arquivo = pasta_desfechos / "desfechos.csv"
    if not arquivo.exists():
        return pd.DataFrame(columns=["chave_caso", *COLUNAS_DESFECHO])
    return pd.read_csv(arquivo, sep=";", dtype=str, encoding="utf-8-sig").fillna("")


def ler_vistos(pasta_desfechos: Path) -> pd.DataFrame:
    arquivo = pasta_desfechos / "casos_vistos.csv"
    if not arquivo.exists():
        return pd.DataFrame(columns=["chave_caso", "primeira_vez"])
    return pd.read_csv(arquivo, sep=";", dtype=str, encoding="utf-8-sig")


def montar_casos(alertas: pd.DataFrame, ids: list[str], desfechos: pd.DataFrame, vistos: pd.DataFrame,
                 lote: dict | None = None) -> pd.DataFrame:
    lote = lote or {"nota_1": (), "sempre": (), "maiores_valores": 10}
    linhas = []
    for chave, grupo in alertas.groupby("chave_caso", sort=False):
        por_regra = grupo.groupby("id_regra")["pontuacao"].max()
        principal = grupo.sort_values(["pontuacao", "valor_envolvido"], ascending=False).iloc[0]
        motivos_lote = []
        if len(por_regra) >= 3:
            motivos_lote.append(f"{len(por_regra)} regras")
        motivos_lote += [f"{r} nota 1" for r in lote["nota_1"] if por_regra.get(r, 0) >= 0.999]
        motivos_lote += [r for r in lote["sempre"] if r in por_regra.index]
        linhas.append({
            "chave_caso": chave,
            "nota_combinada": nota_combinada(list(por_regra)),
            "qt_regras": len(por_regra),
            "regras": " · ".join(r for r in ids if r in por_regra.index),
            **{r: (1 if r in por_regra.index else None) for r in ids},
            "documento": grupo["documento"].iloc[0],
            "fornecedor": grupo["fornecedor"].iloc[0],
            "departamento": grupo["departamento"].iloc[0],
            "maior_valor": float(grupo["valor_envolvido"].max()),
            "qt_alertas": len(grupo),
            "resumo": f"{principal.id_regra}: {principal.evidencia}"[:500],
            "lote": "; ".join(motivos_lote),
        })
    casos = pd.DataFrame(linhas)
    if casos.empty:
        return casos
    casos = casos.sort_values(["nota_combinada", "qt_regras", "maior_valor"], ascending=False).reset_index(drop=True)
    casos.insert(0, "prioridade", range(1, len(casos) + 1))
    candidatos = casos[(casos["lote"] == "") & (casos["qt_regras"] >= 2)].nlargest(lote["maiores_valores"],
                                                                                    "maior_valor")
    casos.loc[candidatos.index, "lote"] = "2 regras, maior valor"
    casos["lote"] = casos["lote"].map(lambda m: f"1 — {m}" if m else "")
    casos.insert(1, "lote", casos.pop("lote"))
    primeira = dict(zip(vistos["chave_caso"], vistos["primeira_vez"])) if len(vistos) else {}
    hoje = date.today().isoformat()
    casos["na_lista_desde"] = casos["chave_caso"].map(lambda c: primeira.get(c, hoje))
    anteriores = desfechos.set_index("chave_caso")[COLUNAS_DESFECHO] if len(desfechos) else None
    for coluna in COLUNAS_DESFECHO:
        casos[coluna] = casos["chave_caso"].map(anteriores[coluna]).fillna("") if anteriores is not None else ""
    casos["situacao"] = casos["situacao"].replace("", "Pendente")
    return casos


def montar_fornecedores(casos: pd.DataFrame) -> pd.DataFrame:
    if casos.empty:
        return pd.DataFrame()
    agrupado = casos.groupby("documento").agg(
        fornecedor=("fornecedor", lambda s: next((n for n in s if n), "")),
        qt_casos=("chave_caso", "size"),
        regras=("regras", lambda s: " · ".join(sorted({r for x in s for r in x.split(" · ") if r}))),
        nota_max=("nota_combinada", "max"),
        valor_total=("maior_valor", "sum"),
    ).reset_index()
    agrupado["qt_regras"] = agrupado["regras"].str.count("·") + 1
    return agrupado.sort_values(["nota_max", "qt_regras", "valor_total"], ascending=False).reset_index(drop=True)


def atualizar_vistos(pasta_desfechos: Path, casos: pd.DataFrame) -> pd.DataFrame:
    vistos = ler_vistos(pasta_desfechos)
    novos = casos.loc[~casos["chave_caso"].isin(vistos["chave_caso"]), ["chave_caso", "na_lista_desde"]]
    vistos = pd.concat([vistos, novos.rename(columns={"na_lista_desde": "primeira_vez"})], ignore_index=True)
    pasta_desfechos.mkdir(parents=True, exist_ok=True)
    vistos.to_csv(pasta_desfechos / "casos_vistos.csv", sep=";", index=False, encoding="utf-8-sig")
    return vistos


def gravar_excel(caminho: Path, casos: pd.DataFrame, alertas: pd.DataFrame, fornecedores: pd.DataFrame,
                 regras: pd.DataFrame, ids: list[str], lote: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fonte = Font(name="Arial", size=10)
    negrito = Font(name="Arial", size=10, bold=True)
    titulo = Font(name="Arial", size=14, bold=True)
    cinza = Font(name="Arial", size=9, color="808080")
    cabecalho = PatternFill("solid", fgColor="D9E1F2")
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    moeda = '"R$" #,##0.00'

    wb = Workbook()

    ws = wb.active
    ws.title = "Casos"
    colunas = ["prioridade", "lote", "nota_combinada", "qt_regras", "regras", *ids, "documento", "fornecedor",
               "departamento", "maior_valor", "qt_alertas", "na_lista_desde", "resumo", *COLUNAS_DESFECHO,
               "chave_caso"]
    titulos = {"prioridade": "Prioridade", "lote": "Lote sugerido", "nota_combinada": "Nota combinada",
               "qt_regras": "Nº de regras", "regras": "Regras", "documento": "CNPJ", "fornecedor": "Fornecedor",
               "departamento": "Departamento", "maior_valor": "Maior valor envolvido (R$)",
               "qt_alertas": "Nº de alertas", "na_lista_desde": "Na lista desde",
               "resumo": "Alerta principal (evidência)", "chave_caso": "chave_caso (não alterar)",
               **TITULOS_DESFECHO}
    letra = {c: get_column_letter(i + 1) for i, c in enumerate(colunas)}
    for j, coluna in enumerate(colunas, start=1):
        celula = ws.cell(row=1, column=j, value=titulos.get(coluna, coluna))
        celula.font = negrito
        celula.fill = amarelo if coluna in COLUNAS_DESFECHO else cabecalho
        celula.alignment = Alignment(wrap_text=True, vertical="center")
    for i, registro in enumerate(casos[colunas].itertuples(index=False), start=2):
        for j, (coluna, valor) in enumerate(zip(colunas, registro), start=1):
            if valor is None or (isinstance(valor, float) and math.isnan(valor)) or valor == "":
                valor = None
            if coluna == "data_analise" and valor:
                try:
                    valor = datetime.fromisoformat(str(valor)[:10]).date()
                except ValueError:
                    pass
            celula = ws.cell(row=i, column=j, value=valor)
            celula.font = cinza if coluna == "chave_caso" else fonte
            if coluna in COLUNAS_DESFECHO:
                celula.fill = amarelo
            if coluna == "maior_valor":
                celula.number_format = moeda
            elif coluna == "nota_combinada":
                celula.number_format = "0.00"
            elif coluna == "data_analise":
                celula.number_format = "dd/mm/yyyy"
    ultima = max(len(casos) + 1, 2)
    larguras = {"prioridade": 9, "lote": 20, "nota_combinada": 10, "qt_regras": 9, "regras": 16, "documento": 17,
                "fornecedor": 34, "departamento": 14, "maior_valor": 18, "qt_alertas": 9, "na_lista_desde": 12,
                "resumo": 70, "situacao": 15, "motivo": 28, "providencia": 28, "analista": 16, "data_analise": 12,
                "observacao": 40, "chave_caso": 30}
    for coluna in colunas:
        ws.column_dimensions[letra[coluna]].width = larguras.get(coluna, 6)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{letra['chave_caso']}{ultima}"
    for coluna, opcoes in (("situacao", SITUACOES), ("motivo", MOTIVOS), ("providencia", PROVIDENCIAS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor fora da lista",
                            error="Escolha um valor da lista (veja a aba Como usar).")
        dv.add(f"{letra[coluna]}2:{letra[coluna]}{max(ultima, 2000)}")
        ws.add_data_validation(dv)
    dv_data = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                             showErrorMessage=True, error="Use uma data (dd/mm/aaaa).")
    dv_data.add(f"{letra['data_analise']}2:{letra['data_analise']}{max(ultima, 2000)}")
    ws.add_data_validation(dv_data)

    def faixa(coluna: str) -> str:
        return f"Casos!${letra[coluna]}$2:${letra[coluna]}${ultima}"

    rs = wb.create_sheet("Resumo", 0)
    rs["A1"] = "Lista única de alertas — resumo"
    rs["A1"].font = titulo
    rs["A2"] = f"Montada em {datetime.now():%d/%m/%Y %H:%M} com o arquivo mais recente de cada regra (tabela abaixo)."
    rs["A4"], rs["B4"], rs["C4"] = "Situação", "Casos", "Maior valor envolvido (R$)"
    for i, situacao in enumerate(SITUACOES, start=5):
        rs[f"A{i}"] = situacao
        rs[f"B{i}"] = f'=COUNTIF({faixa("situacao")},A{i})'
        rs[f"C{i}"] = f'=SUMIF({faixa("situacao")},A{i},{faixa("maior_valor")})'
        rs[f"C{i}"].number_format = moeda
    total = 5 + len(SITUACOES)
    rs[f"A{total}"] = "Total"
    rs[f"B{total}"] = f"=SUM(B5:B{total - 1})"
    rs[f"C{total}"] = f"=SUM(C5:C{total - 1})"
    rs[f"C{total}"].number_format = moeda
    rs[f"A{total + 1}"] = "Lote 1 sugerido"
    rs[f"B{total + 1}"] = f'=COUNTIF({faixa("lote")},"1*")'
    rs[f"A{total + 2}"] = "Lote 1 ainda pendente"
    rs[f"B{total + 2}"] = f'=COUNTIFS({faixa("lote")},"1*",{faixa("situacao")},"Pendente")'

    inicio = total + 4
    cab = ["Regra", "Nome", "Versão", "Arquivo usado", "Alertas", "Casos com a regra", "Procedentes",
           "Improcedentes", "Taxa de acerto"]
    for j, texto in enumerate(cab, start=1):
        rs.cell(row=inicio, column=j, value=texto)
    for i, r in enumerate(regras.itertuples(index=False), start=inicio + 1):
        rs.cell(row=i, column=1, value=r.id_regra)
        rs.cell(row=i, column=2, value=r.nome)
        rs.cell(row=i, column=3, value=r.versao)
        rs.cell(row=i, column=4, value=r.arquivo)
        rs.cell(row=i, column=5, value=int(r.alertas))
        rs.cell(row=i, column=6, value=f"=COUNTIF({faixa(r.id_regra)},1)")
        rs.cell(row=i, column=7, value=f'=COUNTIFS({faixa(r.id_regra)},1,{faixa("situacao")},"Procedente")')
        rs.cell(row=i, column=8, value=f'=COUNTIFS({faixa(r.id_regra)},1,{faixa("situacao")},"Improcedente")')
        rs.cell(row=i, column=9, value=f'=IFERROR(G{i}/(G{i}+H{i}),"-")')
        rs.cell(row=i, column=9).number_format = "0%"
    nota = inicio + len(regras) + 2
    rs[f"A{nota}"] = ("Taxa de acerto = procedentes ÷ (procedentes + improcedentes) entre os casos em que a regra "
                      "apareceu. Mede quais regras acertam mais e orienta o ajuste dos pesos.")
    for linha in rs.iter_rows():
        for celula in linha:
            if celula.value is not None and celula.font != titulo:
                celula.font = negrito if celula.row in (4, inicio, total) else fonte
    for celula in rs[4] + rs[inicio]:
        if celula.value:
            celula.fill = cabecalho
    for coluna, largura in zip("ABCDEFGHI", (16, 52, 8, 34, 9, 12, 12, 13, 12)):
        rs.column_dimensions[coluna].width = largura

    cu = wb.create_sheet("Como usar", 0)
    textos = [
        ("Lista única de alertas — como usar", titulo),
        ("", fonte),
        ("O que é: todos os alertas das regras, agrupados em casos. Um caso é um fornecedor (CNPJ) num "
         "departamento. A aba Casos traz um caso por linha, do mais para o menos prioritário.", fonte),
        ("Nota combinada = 1 − (1 − nota da regra A) × (1 − nota da regra B) × … — cada regra é um indício; "
         "quanto mais regras apontam o mesmo caso, maior a nota (máximo 1).", fonte),
        ("", fonte),
        ("O que preencher: só as colunas AMARELAS da aba Casos (Situação, Motivo, Providência, Analista, Data da "
         "análise, Observação). Não altere as demais nem a coluna chave_caso.", negrito),
        ("Por onde começar: filtre a coluna Lote sugerido (casos com 3 ou mais regras, "
         + ", ".join([f"{r} com nota 1" for r in lote["nota_1"]] + list(lote["sempre"]))
         + f" e os {lote['maiores_valores']} maiores valores com 2 regras). A aba Roteiro diz, para cada regra, "
           "o que verificar e o que costuma explicar o alerta.", fonte),
        ("1. Abra a aba Casos, filtre por lote, departamento, regra ou prioridade e analise o caso.", fonte),
        ("2. Escolha a Situação na lista. Improcedente ou Fora de escopo: diga o Motivo. Procedente: a Providência.",
         fonte),
        ("3. Salve a planilha e registre os desfechos (python -m risco_compras.resultados.lista_unica --registrar). "
         "O que foi concluído vai para a base de desfechos e a próxima lista já vem preenchida.", fonte),
        ("", fonte),
        ("Exemplo de preenchimento (uma linha da aba Casos, só as colunas amarelas):", negrito),
    ]
    for i, (texto, estilo) in enumerate(textos, start=1):
        cu.cell(row=i, column=1, value=texto).font = estilo
    exemplo = ["Improcedente", "Erro de cadastro ou de dado", "", "A. Analista", date(2025, 3, 10),
               "Valor do contrato estava cadastrado como mensal; o total anual confere com o pago."]
    linha_ex = len(textos) + 1
    for j, (coluna, valor) in enumerate(zip(COLUNAS_DESFECHO, exemplo), start=1):
        c1 = cu.cell(row=linha_ex, column=j, value=TITULOS_DESFECHO[coluna])
        c1.font, c1.fill = negrito, amarelo
        c2 = cu.cell(row=linha_ex + 1, column=j, value=valor or None)
        c2.font, c2.fill = fonte, amarelo
        if isinstance(valor, date):
            c2.number_format = "dd/mm/yyyy"
    linha = linha_ex + 3
    for rotulo, opcoes in (("Situação", SITUACOES), ("Motivo", MOTIVOS), ("Providência", PROVIDENCIAS)):
        cu.cell(row=linha, column=1, value=f"Valores de {rotulo}:").font = negrito
        cu.cell(row=linha, column=2, value=" · ".join(opcoes)).font = fonte
        linha += 1
    linha += 1
    cu.cell(row=linha, column=1, value="Regras desta lista:").font = negrito
    for r in regras.itertuples(index=False):
        linha += 1
        cu.cell(row=linha, column=1, value=r.id_regra).font = fonte
        cu.cell(row=linha, column=2, value=f"{r.nome} (v{r.versao})").font = fonte
    cu.column_dimensions["A"].width = 26
    for coluna in "BCDEF":
        cu.column_dimensions[coluna].width = 30

    fs = wb.create_sheet("Fornecedores")
    cab_f = ["CNPJ", "Fornecedor", "Nº de casos", "Regras", "Nº de regras", "Maior nota",
             "Soma dos maiores valores (R$)", "Casos procedentes", "Casos improcedentes"]
    for j, texto in enumerate(cab_f, start=1):
        c = fs.cell(row=1, column=j, value=texto)
        c.font, c.fill = negrito, cabecalho
    for i, f in enumerate(fornecedores.itertuples(index=False), start=2):
        valores = [f.documento, f.fornecedor, int(f.qt_casos), f.regras, int(f.qt_regras), float(f.nota_max),
                   float(f.valor_total),
                   f'=COUNTIFS({faixa("documento")},A{i},{faixa("situacao")},"Procedente")',
                   f'=COUNTIFS({faixa("documento")},A{i},{faixa("situacao")},"Improcedente")']
        for j, valor in enumerate(valores, start=1):
            fs.cell(row=i, column=j, value=valor).font = fonte
        fs.cell(row=i, column=6).number_format = "0.00"
        fs.cell(row=i, column=7).number_format = moeda
    for coluna, largura in zip("ABCDEFGHI", (17, 40, 10, 22, 10, 10, 20, 12, 13)):
        fs.column_dimensions[coluna].width = largura
    fs.freeze_panes = "C2"
    fs.auto_filter.ref = f"A1:I{max(len(fornecedores) + 1, 2)}"

    guia = roteiros()
    ro = wb.create_sheet("Roteiro", wb.sheetnames.index("Casos") + 1)
    cab_r = ["Regra", "Nome", "O que o alerta indica", "O que verificar", "Falsos positivos comuns (e Motivo)"]
    for j, texto in enumerate(cab_r, start=1):
        c = ro.cell(row=1, column=j, value=texto)
        c.font, c.fill = negrito, cabecalho
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(regras.itertuples(index=False), start=2):
        item = guia.get(r.id_regra, {})
        valores = [r.id_regra, f"{r.nome} (v{r.versao})", item.get("indica", ""), item.get("verificar", ""),
                   item.get("falsos_positivos", "")]
        for j, valor in enumerate(valores, start=1):
            c = ro.cell(row=i, column=j, value=valor or None)
            c.font = negrito if j == 1 else fonte
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for coluna, largura in zip("ABCDE", (7, 28, 48, 52, 48)):
        ro.column_dimensions[coluna].width = largura
    ro.freeze_panes = "C2"

    al = wb.create_sheet("Alertas")
    col_a = ["chave_caso", "id_regra", "versao_regra", "pontuacao", "valor_envolvido", "documento", "departamento",
             "tipo_objeto", "id_objeto", "classe", "evidencia", "data_execucao"]
    tit_a = ["chave_caso", "Regra", "Versão", "Pontuação", "Valor envolvido (R$)", "CNPJ", "Departamento",
             "Tipo", "Objeto", "Classe", "Evidência", "Executada em"]
    for j, texto in enumerate(tit_a, start=1):
        c = al.cell(row=1, column=j, value=texto)
        c.font, c.fill = negrito, cabecalho
    for i, registro in enumerate(alertas[col_a].itertuples(index=False), start=2):
        for j, valor in enumerate(registro, start=1):
            al.cell(row=i, column=j, value=valor if valor != "" else None).font = fonte
        al.cell(row=i, column=4).number_format = "0.00"
        al.cell(row=i, column=5).number_format = moeda
    for coluna, largura in zip("ABCDEFGHIJKL", (30, 7, 7, 10, 18, 17, 14, 14, 30, 24, 90, 19)):
        al.column_dimensions[coluna].width = largura
    al.freeze_panes = "C2"
    al.auto_filter.ref = f"A1:L{max(len(alertas) + 1, 2)}"

    wb.active = wb.sheetnames.index("Casos")
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def gravar_painel(pasta_painel: Path, casos: pd.DataFrame, alertas: pd.DataFrame, fornecedores: pd.DataFrame,
                  desfechos: pd.DataFrame) -> None:
    pasta_painel.mkdir(parents=True, exist_ok=True)
    situacao = casos.set_index("chave_caso")["situacao"] if len(casos) else pd.Series(dtype=str)
    alertas = alertas.assign(situacao=alertas["chave_caso"].map(situacao).fillna("Pendente"))
    for nome, tabela in (("casos", casos), ("alertas", alertas), ("fornecedores", fornecedores),
                         ("desfechos", desfechos)):
        tabela.to_csv(pasta_painel / f"{nome}.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig")
        try:
            tabela.astype({c: "string" for c in tabela.columns if tabela[c].dtype == object}).to_parquet(
                pasta_painel / f"{nome}.parquet", index=False)
        except ImportError:
            pass


def roteiros(caminho: Path | None = None) -> dict[str, dict]:
    """Roteiro de análise por regra (config/roteiro_analise.yaml), editável pela equipe."""
    import yaml

    caminho = caminho or RAIZ / "config" / "roteiro_analise.yaml"
    if not caminho.exists():
        return {}
    with open(caminho, encoding="utf-8") as arquivo:
        return yaml.safe_load(arquivo) or {}


def nomes_fornecedores() -> dict[str, str]:
    """CNPJ → razão social, lidos do banco analisado. Sem acesso ao banco, a coluna fica vazia."""
    from sqlalchemy import text

    from ..conexao import engine

    with engine().connect() as conexao:
        tabela = pd.read_sql(text("SELECT cnpj, razao_social FROM fornecedores"), conexao)
    return dict(zip(tabela["cnpj"].astype(str), tabela["razao_social"].astype(str)))


def montar(pasta: Path | None = None, usar_banco: bool = True) -> Path:
    p = pastas(pasta)
    ativas = regras_ativas()
    ids = list(ativas)
    arquivos = ultimos_arquivos(p["base"], ids)
    faltando = [r for r in ids if r not in arquivos]
    if faltando:
        print(f"  aviso: sem arquivo de alertas para {', '.join(faltando)}", flush=True)
    nomes = {}
    if usar_banco:
        try:
            nomes = nomes_fornecedores()
        except Exception as erro:
            print(f"  aviso: nomes dos fornecedores indisponíveis ({type(erro).__name__})", flush=True)
    alertas = carregar_alertas(arquivos, nomes)
    if alertas.empty:
        raise FileNotFoundError(f"Nenhum arquivo alertas_<regra>_<carimbo>.csv em {p['base']}")
    desfechos = ler_desfechos(p["desfechos"])
    ids_presentes = [r for r in ids if r in arquivos]
    lote = config_lote()
    casos = montar_casos(alertas, ids_presentes, desfechos, ler_vistos(p["desfechos"]), lote)
    atualizar_vistos(p["desfechos"], casos)
    fornecedores = montar_fornecedores(casos)
    regras = pd.DataFrame([{"id_regra": r, "nome": ativas[r].get("nome", ""), "versao": str(ativas[r].get("versao", "")),
                            "arquivo": arquivos[r].name,
                            "alertas": int(((alertas.id_regra == r) & ~alertas.espelho).sum())}
                           for r in ids_presentes])
    destino = p["base"] / f"lista_unica_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    gravar_excel(destino, casos, alertas, fornecedores, regras, ids_presentes, lote)
    gravar_painel(p["painel"], casos, alertas, fornecedores, desfechos)
    multi = int((casos.qt_regras > 1).sum())
    ja = int((casos.situacao != "Pendente").sum())
    espelhos = int(alertas["espelho"].sum())
    print(f"{len(alertas) - espelhos} alertas de {len(ids_presentes)} regras"
          f"{f' (+{espelhos} espelhados dos pares da {REGRA_DE_PARES})' if espelhos else ''} -> {len(casos)} casos "
          f"({multi} com mais de uma regra), {len(fornecedores)} fornecedores; "
          f"{ja} caso(s) já com desfecho registrado")
    print(f"Planilha: {destino}")
    print(f"Painel: {p['painel']} (casos, alertas, fornecedores, desfechos)")
    return destino


def lista_mais_recente(pasta: Path) -> Path:
    arquivos = [a for a in sorted(Path(pasta).glob("lista_unica_*.xlsx")) if not a.name.startswith("~$")]
    if not arquivos:
        raise FileNotFoundError(f"Nenhuma lista_unica_*.xlsx em {pasta}")
    return arquivos[-1]


def _texto(valor) -> str:
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    if isinstance(valor, (datetime, date, pd.Timestamp)):
        return valor.strftime("%Y-%m-%d")
    return str(valor).strip()


def ler_desfechos_da_planilha(caminho: Path) -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(caminho, data_only=True, read_only=True)
    if "Casos" not in wb.sheetnames:
        raise ValueError(f"{caminho.name}: aba Casos não encontrada")
    linhas = list(wb["Casos"].iter_rows(values_only=True))
    wb.close()
    titulos = {v: k for k, v in TITULOS_DESFECHO.items()}
    titulos.update({"chave_caso (não alterar)": "chave_caso", "CNPJ": "documento", "Fornecedor": "fornecedor",
                    "Departamento": "departamento", "Regras": "regras", "Nota combinada": "nota_combinada"})
    cabecalho = [titulos.get(str(c), str(c)) for c in linhas[0]]
    faltando = [c for c in ["chave_caso", *COLUNAS_DESFECHO] if c not in cabecalho]
    if faltando:
        raise ValueError(f"{caminho.name}: aba Casos sem as colunas {faltando}")
    tabela = pd.DataFrame(linhas[1:], columns=cabecalho, dtype=object)
    tabela = tabela[[c for c in ["chave_caso", "documento", "fornecedor", "departamento", "regras", "nota_combinada",
                                 *COLUNAS_DESFECHO] if c in tabela.columns]]
    for coluna in tabela.columns:
        tabela[coluna] = tabela[coluna].map(_texto)
    return tabela[tabela["chave_caso"] != ""]


def registrar(caminho: Path, pasta: Path | None = None) -> dict:
    """Leva os desfechos preenchidos na planilha para a base (dados/resultados/desfechos)."""
    p = pastas(pasta)
    planilha = ler_desfechos_da_planilha(Path(caminho))
    preenchidos = planilha[planilha["situacao"].isin(SITUACOES) & (planilha["situacao"] != "Pendente")].copy()
    invalidos = planilha[~planilha["situacao"].isin(SITUACOES) & (planilha["situacao"] != "")]
    sem_motivo = preenchidos[preenchidos["situacao"].isin(["Improcedente", "Fora de escopo"])
                             & (preenchidos["motivo"] == "")]
    base = ler_desfechos(p["desfechos"])
    atual = base.set_index("chave_caso") if len(base) else pd.DataFrame(columns=COLUNAS_DESFECHO)
    agora = datetime.now().isoformat(timespec="seconds")
    novos, alterados, iguais = [], [], 0
    for registro in preenchidos.to_dict("records"):
        chave = registro["chave_caso"]
        if chave in atual.index:
            if all(_texto(atual.at[chave, c]) == registro[c] for c in COLUNAS_DESFECHO):
                iguais += 1
                continue
            alterados.append(chave)
        else:
            novos.append(chave)
        registro.update({"registrado_em": agora, "planilha": Path(caminho).name})
        campos = {k: v for k, v in registro.items() if k != "chave_caso"}
        atual.loc[chave, list(campos)] = pd.Series(campos)
    mudou = novos + alterados
    p["desfechos"].mkdir(parents=True, exist_ok=True)
    if mudou:
        nova_base = atual.reset_index().rename(columns={"index": "chave_caso"}).fillna("")
        ordem = ["chave_caso", "documento", "fornecedor", "departamento", "regras", "nota_combinada",
                 *COLUNAS_DESFECHO, "registrado_em", "planilha"]
        nova_base = nova_base[[c for c in ordem if c in nova_base.columns]]
        nova_base.to_csv(p["desfechos"] / "desfechos.csv", sep=";", index=False, encoding="utf-8-sig")
        historico = preenchidos[preenchidos["chave_caso"].isin(mudou)].assign(registrado_em=agora,
                                                                              planilha=Path(caminho).name)
        arquivo_hist = p["desfechos"] / "desfechos_historico.csv"
        historico.to_csv(arquivo_hist, sep=";", index=False, encoding="utf-8-sig", mode="a",
                         header=not arquivo_hist.exists())
    return {"novos": len(novos), "alterados": len(alterados), "iguais": iguais,
            "invalidos": invalidos[["chave_caso", "situacao"]].to_dict("records"),
            "sem_motivo": len(sem_motivo)}


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Lista única de alertas e registro de desfechos")
    parser.add_argument("--registrar", nargs="?", const="", metavar="PLANILHA",
                        help="Registra os desfechos da planilha (sem caminho: a lista_unica mais recente)")
    parser.add_argument("--pasta", help="Pasta dos alertas (padrão: dados/resultados)")
    parser.add_argument("--sem-banco", action="store_true", help="Não busca os nomes dos fornecedores no banco")
    args = parser.parse_args(argv)
    pasta = Path(args.pasta) if args.pasta else None
    if args.registrar is not None:
        caminho = Path(args.registrar) if args.registrar else lista_mais_recente(pastas(pasta)["base"])
        print(f"Registrando desfechos de {caminho.name}...", flush=True)
        r = registrar(caminho, pasta)
        print(f"{r['novos']} novo(s), {r['alterados']} alterado(s), {r['iguais']} sem mudança.")
        if r["sem_motivo"]:
            print(f"  atenção: {r['sem_motivo']} improcedente(s)/fora de escopo sem Motivo — preencha na próxima lista.")
        for item in r["invalidos"]:
            print(f"  ignorado (Situação fora da lista): {item['chave_caso']} = {item['situacao']!r}")
        print(f"Base: {pastas(pasta)['desfechos'] / 'desfechos.csv'}")
        return
    montar(pasta, usar_banco=not args.sem_banco)


if __name__ == "__main__":
    main()
