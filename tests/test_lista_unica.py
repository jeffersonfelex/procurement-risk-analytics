import pandas as pd
import pytest
from openpyxl import load_workbook

from risco_compras.resultados import lista_unica as lu

CNPJ_A, CNPJ_B, CNPJ_C = "11222333000181", "44555666000199", "77888999000100"


def _alertas(regra: str, linhas: list[dict]) -> pd.DataFrame:
    base = {"id_regra": regra, "tipo_objeto": "pedido", "departamento": "TI", "valor_envolvido": 1000.0,
            "pontuacao": 1.0, "evidencia": "[x] teste", "versao_regra": "1.0", "data_execucao": "2025-01-01T00:00:00",
            "status_tratamento": "pendente"}
    return pd.DataFrame([{**base, **linha} for linha in linhas])


@pytest.fixture
def pasta_alertas(tmp_path):
    _alertas("R1", [{"id_objeto": "g1", "cnpj_fornecedor": CNPJ_A, "pontuacao": 0.5, "valor_envolvido": 30000}]
             ).to_csv(tmp_path / "alertas_R1_20250101_100000.csv", sep=";", index=False)
    _alertas("R1", [{"id_objeto": "velho", "cnpj_fornecedor": CNPJ_C}]
             ).to_csv(tmp_path / "alertas_R1_20240101_100000.csv", sep=";", index=False)   # mais antigo: ignorado
    _alertas("R2", [{"id_objeto": "p1", "cnpj_fornecedor": CNPJ_A, "pontuacao": 0.6}]
             ).to_csv(tmp_path / "alertas_R2_20250101_100000.csv", sep=";", index=False)
    _alertas("R3", [{"id_objeto": "COT-1", "cnpj_fornecedor": CNPJ_A, "pontuacao": 0.7, "tipo_objeto": "cotacao",
                     "evidencia": f"[socio] Vínculo com {CNPJ_B} na cotação COT-1"}]
             ).to_csv(tmp_path / "alertas_R3_20250101_100000.csv", sep=";", index=False)
    return tmp_path


def test_nota_combinada():
    assert lu.nota_combinada([0.5]) == 0.5
    assert lu.nota_combinada([0.5, 0.5]) == 0.75
    assert lu.nota_combinada([1.0, 0.2]) == 1.0
    assert lu.nota_combinada([]) == 0.0


def test_ultimos_arquivos_pega_o_carimbo_mais_recente(pasta_alertas):
    arquivos = lu.ultimos_arquivos(pasta_alertas, ["R1", "R2", "R3"])
    assert arquivos["R1"].name == "alertas_R1_20250101_100000.csv"


def test_casos_combinam_regras_e_espelham_pares(pasta_alertas):
    arquivos = lu.ultimos_arquivos(pasta_alertas, ["R1", "R2", "R3"])
    alertas = lu.carregar_alertas(arquivos, {CNPJ_A: "Empresa A"})
    assert alertas["espelho"].sum() == 1                       # o par da R3 vira caso também para a empresa B
    casos = lu.montar_casos(alertas, ["R1", "R2", "R3"], lu.ler_desfechos(pasta_alertas), pd.DataFrame(
        columns=["chave_caso", "primeira_vez"]), {"nota_1": (), "sempre": ("R2",), "maiores_valores": 10})
    primeiro = casos.iloc[0]
    assert primeiro["documento"] == CNPJ_A and primeiro["qt_regras"] == 3 and primeiro["fornecedor"] == "Empresa A"
    assert primeiro["nota_combinada"] == lu.nota_combinada([0.5, 0.6, 0.7])
    assert primeiro["lote"].startswith("1 — 3 regras") and "R2" in primeiro["lote"]
    assert set(casos["documento"]) == {CNPJ_A, CNPJ_B} and (casos["situacao"] == "Pendente").all()


def test_ciclo_planilha_desfecho_e_proxima_lista(pasta_alertas, monkeypatch):
    monkeypatch.setattr(lu, "regras_ativas", lambda: {r: {"nome": f"Regra {r}", "versao": "1.0"}
                                                      for r in ("R1", "R2", "R3")})
    planilha = lu.montar(pasta_alertas, usar_banco=False)
    wb = load_workbook(planilha)
    assert {"Como usar", "Resumo", "Casos", "Roteiro", "Fornecedores", "Alertas"} <= set(wb.sheetnames)

    ws = wb["Casos"]
    cabecalho = [c.value for c in ws[1]]
    col = {nome: cabecalho.index(nome) + 1 for nome in ("Situação", "Motivo (se improcedente)", "CNPJ")}
    linha_a = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, col["CNPJ"]).value == CNPJ_A)
    ws.cell(linha_a, col["Situação"]).value = "Procedente"
    linha_b = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, col["CNPJ"]).value == CNPJ_B)
    ws.cell(linha_b, col["Situação"]).value = "Talvez"            # fora da lista: ignorado e avisado
    wb.save(planilha)

    resultado = lu.registrar(planilha, pasta_alertas)
    assert (resultado["novos"], resultado["alterados"]) == (1, 0)
    assert [i["situacao"] for i in resultado["invalidos"]] == ["Talvez"]
    assert lu.registrar(planilha, pasta_alertas)["iguais"] == 1    # registrar de novo não duplica

    proxima = lu.montar(pasta_alertas, usar_banco=False)
    casos = pd.read_excel(proxima, sheet_name="Casos", dtype=str)
    assert casos.loc[casos["CNPJ"] == CNPJ_A, "Situação"].item() == "Procedente"
    assert (pasta_alertas / "painel" / "casos.csv").exists()
